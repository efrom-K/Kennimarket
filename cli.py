import argparse
import csv
import os
from datetime import datetime, timedelta, timezone

from leadgen import db
from leadgen.config import settings
from leadgen.pipeline import run_dgis_pipeline, run_telegram_pipeline, run_triage


def cmd_init(_args):
    db.init_db(settings.db_path)
    print(f"БД инициализирована: {settings.db_path}")


def cmd_telegram(args):
    if args.channels_file:
        with open(args.channels_file, encoding="utf-8") as f:
            raw = [line.split("#")[0] for line in f]
    else:
        raw = args.channels.split(",")
    channels = [c.strip() for c in raw if c.strip()]
    run_telegram_pipeline(
        channels, limit_per_query=args.limit, min_confidence=args.min_confidence, max_age_days=args.max_age_days
    )


def cmd_tg_discover(args):
    from leadgen.sources.telegram_source import PROBE_DAYS, discover_chats

    # Выключенные вручную чаты (строка начинается с "#") остаются выключенными,
    # включённые перепроверяются заново вместе с новыми кандидатами.
    disabled, seeds = set(), []
    if os.path.exists(args.out):
        with open(args.out, encoding="utf-8") as f:
            for line in f:
                name = line.lstrip("# ").split("#")[0].strip()
                if name:
                    (disabled.add(name) if line.startswith("#") else seeds.append(name))
    chats = discover_chats(seeds=seeds, min_members=args.min_members, depth=args.depth)
    with open(args.out, "w", encoding="utf-8") as f:
        for c in chats:
            prefix = "# " if c["username"] in disabled else ""
            f.write(f"{prefix}{c['username']}  # {c['kind']}, {c['members']} уч., "
                    f"{c['title']} · запросов за {PROBE_DAYS} дн.: {c['hits']}\n")
    print(f"Чатов с запросами на покупку: {len(chats)} -> {args.out}")


def cmd_dgis(args):
    queries = [q.strip() for q in args.queries.split(",")] if args.queries else None
    run_dgis_pipeline(queries=queries, pages_per_query=args.pages, min_confidence=args.min_confidence)


def cmd_stats(_args):
    db.init_db(settings.db_path)
    with db.get_conn(settings.db_path) as conn:
        total = db.count_leads(conn)
        by_source = conn.execute("SELECT source, COUNT(*) FROM leads GROUP BY source").fetchall()
    print(f"Всего лидов: {total} (цель: {settings.target_leads})")
    for row in by_source:
        print(f"  {row[0]}: {row[1]}")


EXPORT_COLUMNS = [  # (заголовок, ширина в xlsx)
    ("№", 6), ("Кто", 16), ("Тип", 14), ("Дата", 11), ("Имя", 22), ("Telegram", 22), ("Телефон", 15),
    ("Как связаться", 34), ("Что ищет (кратко)", 48), ("Локация", 26), ("Площадь", 12), ("Бюджет", 16),
    ("Текст запроса", 70), ("Черновик первого сообщения", 60), ("Сообщение в чате", 30), ("Комментарий", 24),
]
BUYER_ORDER = {"direct": 0, "unknown": 1, "broker": 2}
BUYER_LABEL = {"direct": "Покупатель", "unknown": "Не ясно", "broker": "Брокер"}
KIND_LABEL = {"warehouse": "Склад / земля", "gab": "ГАБ-инвестор"}


def _export_rows(conn, max_age_days: int) -> list:
    query = ("SELECT l.*, r.raw_text, COALESCE(p.draft_edited, p.draft) AS draft FROM leads l"
             " LEFT JOIN raw_items r ON r.url = l.source_ref LEFT JOIN lead_profiles p ON p.lead_id = l.id"
             " WHERE l.status != 'rejected'")
    params = ()
    if max_age_days:
        # лиды без даты (2GIS) не отсекаем — у них нет понятия "свежести"
        query += " AND (l.posted_at IS NULL OR l.posted_at >= ?)"
        params = ((datetime.now(timezone.utc) - timedelta(days=max_age_days)).isoformat(),)
    rows = conn.execute(query, params).fetchall()
    # сначала прямые покупатели, потом брокеры; внутри — свежие сверху
    rows.sort(key=lambda r: r["posted_at"] or "", reverse=True)
    rows.sort(key=lambda r: BUYER_ORDER.get(r["buyer_type"], 1))
    out = []
    for r in rows:
        user, phone, uid = r["telegram_username"], r["phone"], r["tg_user_id"]
        if user:
            how = f"Написать в Telegram @{user}"
        elif uid:
            how = (f"Нет @username: ответить на сообщение в чате или открыть tg://user?id={uid} "
                   "(сработает после захода в этот чат)")
        else:
            how = "Ответить на сообщение в чате"
        if phone:
            how += f"; звонок / WhatsApp {phone}"
        name = (r["contact_name"] or "").strip(" .") or (f"@{user}" if user else "без имени")
        out.append({
            "buyer": r["buyer_type"] or "unknown",
            "values": [
                r["id"], BUYER_LABEL.get(r["buyer_type"], "Не ясно"), KIND_LABEL.get(r["lead_kind"], r["lead_kind"] or ""),
                (r["posted_at"] or "")[:10], name, f"@{user}" if user else "", phone or "", how,
                r["notes"] or "", r["location"] or "", r["area_sqm"] or "", r["budget"] or "",
                (r["raw_text"] or "").strip(), r["draft"] or "", r["source_ref"] or "", r["comment"] or "",
            ],
            "links": {"Telegram": f"https://t.me/{user}" if user else None, "Сообщение в чате": r["source_ref"]},
        })
    return out


def _write_xlsx(path: str, rows: list) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    fills = {"direct": "E2F0D9", "unknown": "FFF2CC", "broker": "F2F2F2"}  # зелёный / жёлтый / серый
    wb = Workbook()
    ws = wb.active
    ws.title = "Лиды"
    headers = [h for h, _ in EXPORT_COLUMNS]
    ws.append(headers)
    thin = Side(style="thin", color="BFBFBF")
    for i, (_, width) in enumerate(EXPORT_COLUMNS, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[c.column_letter].width = width
    for r in rows:
        ws.append(r["values"])
        n = ws.max_row
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=n, column=i)
            c.font = Font(name="Arial", size=10, bold=h in ("Имя", "Кто"))
            c.fill = PatternFill("solid", fgColor=fills.get(r["buyer"], "FFFFFF"))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            link = r["links"].get(h)
            if link:
                c.hyperlink = link
                c.font = Font(name="Arial", size=10, color="0563C1", underline="single")
    ws.freeze_panes = "F2"  # № … Имя и шапка остаются на месте при прокрутке
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    legend = wb.create_sheet("Пояснения")
    for line in [
        ("Цвет строки", ""), ("Зелёный", "Покупатель — ищет для себя, писать в первую очередь"),
        ("Жёлтый", "Не ясно, сам или для клиента"), ("Серый", "Брокер — ищет для клиента, делится комиссией"),
        ("", ""), ("Колонки", ""), ("Как связаться", "Telegram @username, если есть; иначе — ответ на сообщение в чате. "
         "Телефон есть, только если автор сам его написал или открыл в Telegram"),
        ("Что ищет (кратко)", "Пояснение модели, почему это лид"),
        ("Черновик первого сообщения", "Готовый текст, если среди ваших объектов есть подходящий склад"),
    ]:
        legend.append(line)
    legend.column_dimensions["A"].width = 28
    legend.column_dimensions["B"].width = 90
    for row in legend.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=10, bold=c.column == 1)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for name, color in (("A2", "E2F0D9"), ("A3", "FFF2CC"), ("A4", "F2F2F2")):
        legend[name].fill = PatternFill("solid", fgColor=color)
    wb.save(path)


def cmd_export(args):
    db.init_db(settings.db_path)
    with db.get_conn(settings.db_path) as conn:
        rows = _export_rows(conn, args.max_age_days)
    if not rows:
        print("Лидов пока нет.")
        return
    if args.out.lower().endswith(".xlsx"):
        _write_xlsx(args.out, rows)
    else:
        # ";" — разделитель, который русский Excel открывает по двойному клику
        with open(args.out, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow([h for h, _ in EXPORT_COLUMNS])
            writer.writerows([" ".join(str(v).split()) for v in r["values"]] for r in rows)
    print(f"Экспортировано {len(rows)} лидов -> {args.out}")


def main():
    parser = argparse.ArgumentParser(description="Прототип сбора холодных лидов на покупку складов (Москва/МО)")
    sub = parser.add_subparsers(dest="command", required=True)

    sub.add_parser("init-db").set_defaults(func=cmd_init)

    p_disc = sub.add_parser("tg-discover", help="Найти публичные Telegram-чаты/каналы по тематическим запросам")
    p_disc.add_argument("--out", default="channels.txt")
    p_disc.add_argument("--min-members", type=int, default=300)
    p_disc.add_argument("--depth", type=int, default=1, help="Сколько раз раскручивать через «похожие каналы»")
    p_disc.set_defaults(func=cmd_tg_discover)

    p_tg = sub.add_parser("telegram", help="Искать buy-intent сигналы в публичных Telegram-каналах и группах")
    src = p_tg.add_mutually_exclusive_group(required=True)
    src.add_argument("--channels", help="Список @username через запятую")
    src.add_argument("--channels-file", help="Файл со списком username (по одному в строке, # — комментарий)")
    p_tg.add_argument("--limit", type=int, default=200, help="Максимум результатов на один поисковый запрос в чате")
    p_tg.add_argument("--max-age-days", type=int, default=90, help="Брать сообщения не старше N дней (0 — вся история)")
    p_tg.add_argument("--min-confidence", type=float, default=0.5)
    p_tg.set_defaults(func=cmd_telegram)

    p_dg = sub.add_parser("dgis", help="Собрать компании-таргеты через 2GIS Catalog API")
    p_dg.add_argument("--queries", default=None, help="Поисковые фразы через запятую (по умолчанию — набор из sources/dgis_source.py)")
    p_dg.add_argument("--pages", type=int, default=5, help="Сколько страниц выдачи на каждый запрос")
    p_dg.add_argument("--min-confidence", type=float, default=0.4)
    p_dg.set_defaults(func=cmd_dgis)

    sub.add_parser("triage", help="Подготовить первое сообщение лидам без черновика").set_defaults(
        func=lambda _a: run_triage())

    sub.add_parser("stats").set_defaults(func=cmd_stats)

    p_exp = sub.add_parser("export")
    p_exp.add_argument("--out", default="leads.xlsx", help=".xlsx — с цветом и ссылками, .csv — плоская таблица")
    p_exp.add_argument("--max-age-days", type=int, default=0, help="Только сообщения не старше N дней (0 — все)")
    p_exp.set_defaults(func=cmd_export)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
